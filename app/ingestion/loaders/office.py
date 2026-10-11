import os
import logfire
import docx
import pptx


def parse_docx(file_path: str) -> str:
    """Extract all text paragraphs and tables from a Word document."""
    doc = docx.Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def parse_pptx(file_path: str) -> str:
    """Extract all text shapes from a PowerPoint presentation."""
    prs = pptx.Presentation(file_path)
    slides_text = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_parts.append(text)
        if slide_parts:
            slides_text.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_parts))
    return "\n\n".join(slides_text)


def parse_office(file_path: str) -> str:
    """
    Parses Office documents (.docx, .pptx) directly using python-docx and python-pptx.
    Runs 100% locally with zero heavy external C++ dependencies or PyTorch.
    """
    with logfire.span("📄 Office Document Parsing", filename=file_path):
        try:
            ext = file_path.lower().rsplit(".", 1)[-1]
            if ext == "docx":
                full_text = parse_docx(file_path)
            elif ext == "pptx":
                full_text = parse_pptx(file_path)
            else:
                raise ValueError(f"Unsupported office file type: {ext}")

            if not full_text.strip():
                logfire.warning(f"⚠️ Empty text extracted from {file_path}")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters from {os.path.basename(file_path)}")

            return full_text
        except Exception as e:
            logfire.error(f"❌ Office Parse Failed: {e}")
            raise e
